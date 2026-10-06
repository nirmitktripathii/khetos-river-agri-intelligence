"""The flow-history tables as one Excel workbook laid out for reading (needs openpyxl).

Imported by build_flow_history.py, which puts the repository root and this folder on the import path. Every sheet
holds plain values except Calculators, whose yellow cells feed live formulas. The wording of the Read me, Limits
and Dictionary sheets lives here; every number in them is taken from the tables.
"""
import math
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pull_rain import SHARE
from src import flow

APP = "https://khetos-river-agri-intelligence.streamlit.app/"
SHORT = {"above": "Above the city", "entering": "Entering the city", "khajuria": "Khajuriya ghat",
         "below": "Below the city", "mouth": "Whole river, at the Ramganga", "ramganga": "Ramganga at Chaubari"}
REACH = {"upper": "Upper (Bhojipura side)", "urban": "Urban (Dohra Road to Bisalpur Road)",
         "lower": "Lower (towards the Ramganga)"}
BASE = flow.BASELINE
BASE_TEXT = f"{BASE[0]} to {BASE[1]}"
FIRST = flow.WARMUP_YEARS[-1] + 1  # first year used in statistics
MOUTH_KM2 = flow.SEGMENTS["mouth"]["area_km2"]
WIDTH_NOISE_M = 0.5  # spread of the unmixed width over land away from the river (unmix_probe.py)
WEIGHTS = " : ".join(f"{v:.2f}" for v in SHARE.values())

BANNER = ("MODELLED, NOT MEASURED. Nobody gauges the Nakatiya: every flow here comes from a weather-driven model "
          "and has not been checked against a measurement on this river.")
ABOUT = {
    "Charts": "Six pictures of the main results.",
    "River points": "The five points the flows are given for, with the average-year figures of each.",
    "Annual": "One row per calendar year: mean flow, volume, highest day and lowest week at every point, with "
              "the year's rain.",
    "Water year": "The same by June-to-May water year, the year in which Indian river records are kept.",
    "Seasons": "One row per season of every year: winter, pre-monsoon, monsoon and post-monsoon.",
    "Months": "One row per month, from January 1940.",
    "Days": "One row per day: the daily flow at every point, with the day's rain and evaporative demand.",
    "Average year": f"The average year {BASE_TEXT}, month by month.",
    "Dependable": "The volume that is met or passed in 50, 75 and 90 % of years, for the year and each season.",
    "Flow duration": "The flow that is equalled or exceeded on a given share of days.",
    "Trends": "Whether each yearly series rises or falls, and whether chance alone could explain it.",
    "Decades": "Ten-year averages of rain, volume, runoff and dry-season flow.",
    "Rain and runoff": "Year by year: the rain, the modelled runoff, and the storm runoff the curve-number method "
                       "gives for three land covers.",
    "Land-cover scenario": "How much more storm runoff the same rain gives as land is built over, by the method "
                           "of Dr. S. S. Tripathi's thesis.",
    "Satellite width": "Open-water width of three reaches on clear Sentinel-2 dates, 2018 to 2025.",
    "Checks": "What the model can be checked against, and what the checks say.",
    "Calculators": "Live formulas: discharge by the float method and by Manning's formula, one day's storm "
                   "runoff, and unit conversions.",
    "Limits": "What these numbers cannot tell you. Read this before quoting any of them.",
    "Dictionary": "Every column explained, and the terms used.",
    "Chart data": "The numbers behind the charts, copied from the other sheets.",
}

FLOW, LOW, VOL, MM, PCT, PVAL, INT = "#,##0.00", "0.000", "#,##0.0", "#,##0", "0.0", "0.0000", "0"
ONE, TWO, DAY, MONTH, TEXT = "0.0", "0.00", "yyyy-mm-dd", "yyyy-mm", None

NAVY, PALE, YELLOW, RED, SLATE, AMBER, GREEN = "1F4E79", "DDEBF7", "FFF2CC", "C00000", "7F7F7F", "BF8F00", "548235"
BLUES = ("9DC3E6", "5B9BD5", "2E75B6", NAVY)
TITLE, HEAD, BOLD = Font(size=14, bold=True, color=NAVY), Font(bold=True, color="FFFFFF"), Font(bold=True)
SECTION, WARN, QUIET = Font(size=12, bold=True, color=NAVY), Font(bold=True, color=RED), Font(color="595959")
FILL_HEAD, FILL_BAND, FILL_INPUT = (PatternFill("solid", fgColor=c) for c in (NAVY, PALE, YELLOW))
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)

# Column descriptions shared by several sheets: (key, heading, number format, meaning).
RAIN = ("rain_mm", "Rain (mm)", MM, f"ERA5 rain over the Nakatiya catchment: three grid cells, north to south, "
        f"weighted {WEIGHTS} by their share of the catchment.")
PET = ("et0_mm", "Evaporative demand (mm)", MM, "Reference evapotranspiration (FAO-56 Penman-Monteith) from ERA5 "
       "weather: what a short, well-watered grass cover would evaporate. Weighted like the rain.")
MEAN = ("mean_m3s", "Mean flow (m³/s)", FLOW, "Average of the daily flows of the period, in cubic metres a second.")
VOLUME = ("volume_mcm", "Volume (million m³)", VOL, "Water that passed the point in the period: the daily flows "
          "added up, each multiplied by the 86,400 seconds of a day.")
PEAK = ("peak_day_m3s", "Highest day (m³/s)", FLOW, "Highest daily-mean flow of the period. The flood peak inside "
        "that day was higher.")
LOW7 = ("low_7day_m3s", "Lowest 7 days (m³/s)", LOW, "Lowest mean over seven days in a row: the usual measure of "
        "how low a river falls.")
DAYS = ("days", "Days of data", INT, "Days of modelled flow in the period.")
WHOLE = ("complete", "Whole period?", TEXT, "\"no\" where days are missing: the totals are then part totals, and "
         "the period is left out of every statistic.")
NOTE = ("note", "Note", TEXT, "Why a year is incomplete or is left out of the statistics.")
BALANCE = ("Whole river: water balance", [
    ("mouth_runoff_mm", "Runoff depth (mm)", MM, f"The volume at the mouth spread evenly over the {MOUTH_KM2:g} km² "
     "catchment, so that it can be set beside the rain."),
    ("mouth_runoff_pct_of_rain", "Runoff as % of rain", PCT, "Runoff depth divided by the rain of the same period. "
     "Blank for a part year."),
    ("mouth_volume_10yr_mcm", "Mean volume of 10 years (million m³)", VOL, "Mean volume of the ten years ending "
     f"with this one, from {FIRST}.")])

TERMS = [
    ("m³/s (cubic metre a second, \"cumec\")", "Flow: the volume of water passing a point every second. 1 m³/s is "
     "1,000 litres a second, 35.3 cusec (cubic feet a second) or 86.4 million litres a day."),
    ("million m³ (MCM)", "Volume. 1 million m³ is 100 hectare-metres or 100 crore litres; a flow of 1 m³/s kept up "
     "for a day is 0.0864 million m³, and for a year 31.5 million m³."),
    ("Runoff depth (mm)", f"A volume spread evenly over the catchment, so that it can be compared with rain. 1 mm "
     f"over the {MOUTH_KM2:g} km² of the Nakatiya catchment is {MOUTH_KM2 / 1000:.4f} million m³."),
    ("Catchment", "The land that drains to a point on the river."),
    ("Daily-mean flow", "The average flow over one day. A flood that lasts a few hours shows as a much lower daily "
     "mean than its peak."),
    ("Water year", "June to May, so that one monsoon and the dry season it feeds fall in the same year. Written "
     "1985-86."),
    ("Season", "India Meteorological Department seasons: winter January-February, pre-monsoon March-May, monsoon "
     "June-September, post-monsoon October-December."),
    ("Dependable volume", "The volume met or passed in a stated share of years: \"75 % dependable\" is reached in "
     "three years out of four. Years are ranked and given the Weibull position m / (n + 1)."),
    ("Flow-duration curve", "For every flow, the share of days on which the river carries at least that much."),
    ("Reanalysis (ERA5)", "Past weather reconstructed by running a weather model over the past and steering it "
     "with the observations of the time. ERA5 is the reanalysis of the European Centre for Medium-Range Weather "
     "Forecasts (ECMWF), from 1940."),
    ("GEOGLOWS", "Group on Earth Observations Global Water Sustainability: a global river model that routes the "
     "runoff of ERA5 down a river network mapped from satellite elevation data (TDX-Hydro)."),
    ("Curve number (CN)", "The number, from about 30 to 100, by which the Soil Conservation Service method turns a "
     "day's rain into storm runoff: higher for sealed or clayey ground. CN-II is for average soil moisture; the "
     "rain of the five days before moves a day to CN-I (dry soil) or CN-III (wet soil)."),
    ("Theil-Sen slope", "The median of the slopes between every pair of years: a trend line that a few extreme "
     "years cannot pull."),
    ("Mann-Kendall test, p", "A test of whether a series rises or falls more steadily than chance would give. p is "
     "the chance of so steady a run where there is no real trend; below 0.05 is the usual bar."),
    ("p allowing for persistence", "The same test after allowing for wet years following wet years, which makes "
     "chance runs longer (Yue and Wang, 2004). The reading in words uses this value."),
    ("Sub-pixel unmixing", "Working out what share of a 10 m satellite pixel is water from its colour, so that a "
     "stream narrower than one pixel can still be measured as an area."),
    ("Float method", "Discharge from a tape, a float and a stopwatch: cross-section x surface speed x a "
     "coefficient near 0.85, because water at the surface runs faster than the average."),
    ("Manning's formula", "Mean speed of water in a channel from its shape, its slope and the roughness of its bed: "
     "V = R^(2/3) x S^(1/2) / n, with R the area divided by the wetted perimeter."),
]


def band(key):
    return f"{SHORT[key]} ({flow.SEGMENTS[key]['area_km2']:,.0f} km²)"


def seg_blocks(*stats):
    """One block of columns for each river point."""
    return [(band(k), [(f"{k}_{s}", label, fmt, meaning) for s, label, fmt, meaning in stats]) for k in flow.SEGMENTS]


def clean(v):
    """A value openpyxl can write: blank for NaN, yes / no for a boolean, plain Python numbers and dates."""
    if v is None or v is pd.NaT or (isinstance(v, (float, np.floating)) and not math.isfinite(v)):
        return None
    if isinstance(v, (bool, np.bool_)):
        return "yes" if v else "no"
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, pd.Timestamp):
        return v.date()
    return v


def say(row):
    """A trend-table row in words: "a clear rise (+23.3 % per decade)"."""
    word = {"no trend detected": "no trend", "too few years": "too few years to say"}.get(row["reading"])
    tail = ", within what chance gives" if row["reading"] == "no trend detected" else ""
    return f"{word or 'a ' + row['reading']} ({row['pct_per_decade']:+.1f} % per decade{tail})"


class Panel:
    """The rows of one calculator: yellow cells to type in, and formulas that read them by name."""

    def __init__(self, ws, row):
        self.ws, self.row, self.at = ws, row, {}

    def heading(self, text, about):
        self.ws.cell(self.row, 1, text).font = SECTION
        self.ws.cell(self.row + 1, 1, about).font = QUIET
        for col, label in enumerate(("Value", "Unit", "The sample gives"), 2):
            self.ws.cell(self.row + 2, col, label).font = QUIET
        self.row += 3

    def enter(self, name, label, value, unit=""):
        self.ws.cell(self.row, 1, label)
        self.ws.cell(self.row, 2, value).fill = FILL_INPUT
        self.ws.cell(self.row, 3, unit)
        self.at[name] = f"$B${self.row}"
        self.row += 1

    def enter_many(self, name, label, values, slots=10):
        self.ws.cell(self.row, 1, label)
        for i in range(slots):
            self.ws.cell(self.row, 2 + i, values[i] if i < len(values) else None).fill = FILL_INPUT
        self.at[name] = f"$B${self.row}:${get_column_letter(1 + slots)}${self.row}"
        self.row += 1

    def show(self, name, label, formula, unit, sample, fmt=FLOW):
        self.ws.cell(self.row, 1, label)
        cell = self.ws.cell(self.row, 2, "=" + formula.format(**self.at))
        cell.font, cell.number_format = BOLD, fmt or "General"
        self.ws.cell(self.row, 3, unit)
        check = self.ws.cell(self.row, 4, sample)
        check.font, check.number_format = QUIET, fmt or "General"
        self.at[name] = f"$B${self.row}"
        self.row += 1

    def line(self, text, font=None):
        cell = self.ws.cell(self.row, 1, text)
        if font:
            cell.font = font
        self.row += 1


class Writer:
    """Builds the workbook sheet by sheet, collecting column widths and the column dictionary on the way."""

    def __init__(self, t):
        self.t = t
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.columns, self.widths, self.notes = [], {}, []
        self.trends = t["trends"].set_index(["segment", "series", t["trends"]["period"].str[:4].astype(int)])

    # ------------------------------------------------------------------------------------ plumbing

    def sheet(self, name, title, tab=NAVY, banner=True):
        ws = self.wb.create_sheet(name)
        ws.sheet_properties.tabColor = tab
        ws["A1"], ws["A2"] = title, ABOUT.get(name)
        ws["A1"].font, ws["A2"].font = TITLE, QUIET
        if banner:
            ws["A3"] = BANNER
            ws["A3"].font = WARN
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        return ws

    def widen(self, ws, col, width):
        self.widths[ws.title, col] = max(self.widths.get((ws.title, col), 0), width)

    def trend(self, segment, series, start=flow.TREND_START):
        return self.trends.loc[(segment, series, start)]

    def table(self, ws, df, blocks, top, freeze=None, filtered=False, describe=True):
        """Write `df` at row `top` under a band row (where any block is named) and a heading row. `blocks` is
        [(band, [(key, heading, number format, meaning), ...]), ...]. Returns the row after the table."""
        banded = any(name for name, _ in blocks)
        head, col, keys, formats, lines = top + banded, 1, [], [], 2
        for name, cols in blocks:
            if banded:
                for c in range(col, col + len(cols)):
                    cell = ws.cell(top, c, name if c == col and name else None)
                    cell.font, cell.fill, cell.alignment = BOLD, FILL_BAND, CENTRE
                if name and len(cols) > 1:
                    ws.merge_cells(start_row=top, start_column=col, end_row=top, end_column=col + len(cols) - 1)
            for key, label, fmt, meaning in cols:
                cell = ws.cell(head, col, label)
                cell.font, cell.fill, cell.alignment = HEAD, FILL_HEAD, CENTRE
                if fmt is TEXT:
                    longest = max((len(str(v)) for v in df[key] if clean(v) is not None), default=0)
                    width = min(max(longest + 2, 9), 62)
                else:
                    width = 11.5 if fmt in (DAY, FLOW, VOL) else 10
                width = max(width, max(len(word) for word in label.split()) + 2)  # no heading breaks a word
                self.widen(ws, col, width)
                lines = max(lines, math.ceil(len(label) / (width - 2)) + 1)
                keys.append(key)
                formats.append(fmt)
                if describe:
                    self.columns.append((ws.title, label, meaning))
                col += 1
        ws.row_dimensions[head].height = 15 * min(lines, 6) + 2
        for r, row in enumerate(df[keys].itertuples(index=False, name=None), head + 1):
            for c, (v, fmt) in enumerate(zip(row, formats), 1):
                v = clean(v)
                if v is None:
                    continue
                cell = ws.cell(r, c, v)
                if fmt and not isinstance(v, str):
                    cell.number_format = fmt
        last = head + len(df)
        if filtered:
            ws.auto_filter.ref = f"A{head}:{get_column_letter(len(keys))}{last}"
        if freeze is not None:
            ws.freeze_panes = ws.cell(head + 1, freeze + 1)
            ws.print_title_rows = f"{top}:{head}"
        return last + 2

    def add(self, name, title, df, blocks, freeze=1, tab=NAVY):
        ws = self.sheet(name, title, tab)
        self.table(ws, df, blocks, top=5, freeze=freeze, filtered=True)
        return ws

    def note(self, ws, row, text, cols=10, font=None):
        """A paragraph merged across `cols` columns. Excel does not fit the height of merged cells to their text,
        so finish() sets it once the column widths are known."""
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
        cell = ws.cell(row, 1, text)
        cell.alignment = WRAP
        if font:
            cell.font = font
        self.notes.append((ws.title, row, len(text), cols))
        return row + 1

    def text_sheet(self, name, title, items, tab=SLATE):
        """A sheet of headings ("h"), paragraphs ("p") and bullets ("b") in one wide, wrapped column."""
        ws = self.sheet(name, title, tab)
        ws["A3"].alignment = WRAP
        self.widths[name, 1] = 125
        for r, (kind, text) in enumerate(items, 5):
            cell = ws.cell(r, 1, ("•  " if kind == "b" else "") + text if text else None)
            cell.alignment = WRAP
            if kind == "h":
                cell.font = SECTION
        return ws

    def finish(self):
        for (name, col), width in self.widths.items():
            self.wb[name].column_dimensions[get_column_letter(col)].width = width
        for name, row, length, cols in self.notes:
            across = sum(self.widths.get((name, c), 8.43) for c in range(1, cols + 1))
            self.wb[name].row_dimensions[row].height = 15 * math.ceil(length / across) + 4

    # -------------------------------------------------------------------------------------- sheets

    def read_me(self):
        t, info = self.t, flow.snapshot_info()
        m = t["segments"].set_index("segment").loc["mouth"]
        ann = t["annual"][t["annual"]["complete"]]
        last_year = int(ann["year"].max())
        sc = t["scenario"].set_index("cn2")
        cn79, cn09 = (t["thesis_years"].set_index("year")["thesis_cn2"][y] for y in (1979, 2009))
        year85, year50 = self.trend("mouth", "year"), self.trend("mouth", "year", 1950)
        rain85, rain50 = self.trend("weather", "rain of the year"), self.trend("weather", "rain of the year", 1950)
        dry, low7 = self.trend("mouth", "pre-monsoon (Mar-May)"), self.trend("mouth", "lowest 7-day flow")
        dry_rain = self.trend("weather", "rain of the dry months (Jan-May)")
        demand = self.trend("weather", "evaporative demand of the year")
        items = [
            ("h", "What this workbook is"),
            ("p", f"A history of how much water the Nakatiya river carries, day by day from 1 January 1940 to "
                  f"{pd.Timestamp(info['last']):%d %B %Y}, at four points along the river and, for comparison, in "
                  f"the Ramganga at Chaubari (Bareilly). It belongs to the KhetOS river observatory: {APP}"),
            ("p", "Nobody measures the flow of the Nakatiya. The flows here come from GEOGLOWS, a global river "
                  "model that turns the rain and evaporation of a weather reconstruction (ERA5) into river flow. "
                  "The model knows nothing about the city: no sewage, no drains, no pumping, no canals, no paved "
                  "ground. So these numbers are what the weather alone would put in the river. They are the "
                  "yardstick a real measurement should be set against, not a substitute for one."),
            ("p", "Beside the modelled flows the workbook holds three things worked out independently: the storm "
                  "runoff that the curve-number method of Dr. S. S. Tripathi's thesis gives for the same rain, a "
                  "long-run water balance, and the open-water width of the river seen by satellite. The "
                  "Calculators sheet turns a tape, a float and a stopwatch into a measured flow."),
            ("", ""),
            ("h", f"The main numbers (average year {BASE_TEXT}, whole river where it meets the Ramganga)"),
            ("b", f"Mean flow {m['mean_m3s']:.2f} m³/s (cubic metres a second). On half of all days the river "
                  f"carries less than {m['median_m3s']:.2f} m³/s."),
            ("b", f"Volume of an average year: {m['volume_mcm']:.0f} million m³. Spread over the {MOUTH_KM2:g} km² "
                  f"catchment that is {m['runoff_mm']:.0f} mm of water, "
                  f"{t['yield_estimates']['pct_of_rain'][0]:.0f} % of the rain."),
            ("b", f"In three years out of four the river carries at least {m['dep75_volume_mcm']:.0f} million m³ "
                  f"(the \"75 % dependable\" volume); in nine out of ten, at least {m['dep90_volume_mcm']:.0f}."),
            ("b", f"{m['monsoon_share_pct']:.0f} % of the water passes in the four monsoon months, June to "
                  f"September. March to May average {m['pre-monsoon_mean_m3s']:.2f} m³/s, and the lowest week of "
                  f"a year averages {m['low_7day_m3s']:.2f} m³/s."),
            ("b", f"Highest daily flow in the record: {m['peak_day_m3s']:.0f} m³/s on "
                  f"{pd.Timestamp(m['peak_day_date']):%d %B %Y}."),
            ("", ""),
            ("h", "What the record says about change"),
            ("b", f"Since {flow.TREND_START}, the first year of the app's satellite record, the modelled yearly "
                  f"flow shows {say(year85)} and the rain shows {say(rain85)}."),
            ("b", f"Since 1950 the modelled yearly flow shows {say(year50)}, and so does the rain: {say(rain50)}. "
                  "This is a change in the weather of the reconstruction, not in land use, and the early decades "
                  "rest on few observations: check it against the rain grids of the India Meteorological "
                  "Department before quoting it."),
            ("b", f"The modelled dry-season flow rises since {flow.TREND_START}: March to May shows {say(dry)}, "
                  f"the lowest week of the year {say(low7)}. The rain of the dry months shows {say(dry_rain)}; "
                  f"what changes is the evaporative demand of the year: {say(demand)}. Less evaporation leaves "
                  "more water in the model's soil to drain to the river."),
            ("b", "Why this matters for the question \"has the river lost water because its banks were built "
                  f"over?\": in this model the weather gives no reason for the dry-season river to have shrunk "
                  f"since {flow.TREND_START}. If the real river has shrunk in the dry months, the weather does not "
                  "explain it and the cause has to be looked for on the ground: pumping, seepage lost under "
                  "built-up banks, drains led elsewhere. Only measured flow can show that the river has shrunk."),
            ("b", f"Land cover and storm runoff: for the same {BASE_TEXT} rain, the rise of the curve number from "
                  f"{cn79} (1979) to {cn09} (2009) found in the thesis raises storm runoff by "
                  f"{sc['change_pct'][cn09]:.0f} %, or {sc['change_mcm'][cn09]:.1f} million m³ a year over the "
                  "Nakatiya catchment. That water reaches the river in hours instead of soaking into the ground."),
            ("", ""),
            ("h", "The sheets"),
            *[("b", f"{name}: {text}") for name, text in ABOUT.items()],
            ("", ""),
            ("h", "Conventions"),
            ("b", "Flows are daily means in m³/s; volumes are in million m³ (1 million m³ = 100 hectare-metres = "
                  "100 crore litres)."),
            ("b", "A year is the calendar year unless the sheet says water year (June to May, written 1985-86). "
                  "Seasons are those of the India Meteorological Department: winter January-February, pre-monsoon "
                  "March-May, monsoon June-September, post-monsoon October-December."),
            ("b", f"\"Average year\" means {BASE_TEXT}, the current 30-year climate normal period. Long-record "
                  f"figures use {FIRST} to {last_year}."),
            ("b", "1940 and 1941 are shown but left out of every statistic, as a warm-up precaution for the first "
                  "two years of the model run. The Checks sheet shows that nothing marks them as wrong."),
            ("b", "A blank cell means no value. \"no\" under \"Whole period?\" marks a year or season with days "
                  "missing."),
            ("b", "Days of flow are days of Coordinated Universal Time (UTC), 05:30 to 05:30 India Standard Time "
                  "(IST); days of rain are India Standard Time days."),
            ("", ""),
            ("h", "Sources and licences"),
            ("b", "River flow: GEOGLOWS (Group on Earth Observations Global Water Sustainability) River Forecast "
                  "System version 2, retrospective simulation. ERA5 runoff routed down the TDX-Hydro river "
                  f"network. Licence CC BY 4.0. {info['api']} Retrieved {info['retrieved']}."),
            ("b", "Weather: ERA5, the fifth-generation reanalysis of the European Centre for Medium-Range Weather "
                  "Forecasts (ECMWF); Hersbach and others (2020), Quarterly Journal of the Royal Meteorological "
                  "Society 146, 1999-2049. Read through the Open-Meteo historical weather service (CC BY 4.0), "
                  "asking for ERA5 only so that the series does not change model part-way."),
            ("b", "Curve numbers: S. S. Tripathi (2017), \"Modelling the Effects of Land Use and Land Cover Change "
                  "on Hydrologic Regime of Ramganga River Basin\", Doctor of Philosophy thesis, Sam Higginbottom "
                  "University of Agriculture, Technology and Sciences, Allahabad. SCS is the Soil Conservation "
                  "Service of the United States, which devised the method."),
            ("b", "Check at Chaubari: INRM Consultants for WWF-India (2015), \"Hydrological Modelling of the "
                  "Ramganga River Basin\": a SWAT (Soil and Water Assessment Tool) model calibrated to the "
                  "Central Water Commission (CWC) gauge."),
            ("b", "Satellite width: Copernicus Sentinel-2 Level-2A images of the European Space Agency, read from "
                  "the Microsoft Planetary Computer."),
            ("b", "Channel roughness: V. T. Chow (1959), Open-Channel Hydraulics. Trend test: Yue and Wang (2004), "
                  "Water Resources Management 18, 201-218."),
            ("", ""),
            ("h", "How to cite"),
            ("p", f"KhetOS river observatory ({date.today():%Y}). Nakatiya river: modelled flow and volume, 1940 "
                  f"to {pd.Timestamp(info['last']).year}. Built from GEOGLOWS version 2 (CC BY 4.0) and ERA5. "
                  f"{APP}"),
            ("p", f"Built on {date.today():%d %B %Y} by research/river_flow/build_flow_history.py. The same code "
                  "(src/flow.py) computes the figures of the app's River Water Watch page, so the two agree."),
        ]
        self.text_sheet("Read me", f"Nakatiya river: modelled flow and volume, 1940 to "
                                   f"{pd.Timestamp(info['last']).year}", items)

    def points(self):
        ws = self.sheet("River points", "The five points the flows are given for")
        s = self.t["segments"].set_index("segment")
        for col, label in enumerate(["Quantity", "Unit", *[SHORT[k] for k in s.index]], 1):
            cell = ws.cell(5, col, label)
            cell.font, cell.fill, cell.alignment = HEAD, FILL_HEAD, CENTRE
            self.widen(ws, col, 58 if col == 1 else 24 if col == 2 else 17)
        ws.row_dimensions[5].height = 32
        rows = [
            ("Name", "", "name", TEXT), ("GEOGLOWS river number", "", "river_id", INT),
            ("Stretch of the mapped river the number stands for", "km from the mapped head", "main_stem_km", TEXT),
            ("Catchment area above the point", "km²", "area_km2", VOL),
            ("Latitude of the point", "degrees north", "outlet_lat", PVAL),
            ("Longitude of the point", "degrees east", "outlet_lon", PVAL),
            (f"Average year, {BASE_TEXT}",),
            ("Mean flow", "m³/s", "mean_m3s", FLOW),
            ("Median daily flow (half of all days are below it)", "m³/s", "median_m3s", FLOW),
            ("Volume of an average year", "million m³", "volume_mcm", VOL),
            ("The same as a depth of water over the catchment", "mm", "runoff_mm", MM),
            ("Volume met or passed in 50 % of years", "million m³", "dep50_volume_mcm", VOL),
            ("Volume met or passed in 75 % of years", "million m³", "dep75_volume_mcm", VOL),
            ("Volume met or passed in 90 % of years", "million m³", "dep90_volume_mcm", VOL),
            *[(f"Mean flow, {name} ({flow.SEASON_MONTHS[name]})", "m³/s", f"{name}_mean_m3s", FLOW)
              for name in flow.SEASONS],
            ("Share of the year's volume that passes in the monsoon", "%", "monsoon_share_pct", PCT),
            ("Lowest 7-day flow of a year, averaged over the years", "m³/s", "low_7day_m3s", LOW),
            ("Days with less than 0.1 m³/s", "% of days", "days_below_0p1_pct", PCT),
            (f"Whole record, from {FIRST}",),
            ("Highest daily flow", "m³/s", "peak_day_m3s", FLOW),
            ("Date of the highest daily flow", "", "peak_day_date", DAY),
        ]
        for r, row in enumerate(rows, 6):
            if len(row) == 1:
                for col in range(1, len(s) + 3):
                    cell = ws.cell(r, col, row[0] if col == 1 else None)
                    cell.font, cell.fill = BOLD, FILL_BAND
                continue
            label, unit, key, fmt = row
            for col, text in ((1, label), (2, unit)):
                ws.cell(r, col, text).alignment = Alignment(vertical="top")
            for col, k in enumerate(s.index, 3):
                cell = ws.cell(r, col, clean(s[key][k]))
                cell.alignment = Alignment(horizontal="right", wrap_text=key == "name", vertical="top")
                if fmt:
                    cell.number_format = fmt
        ws.freeze_panes = "C6"
        self.note(ws, 6 + len(rows) + 1, "A flow is the flow leaving a model river segment, so each point is the "
                  "downstream end of its segment. The model's river starts 25 to 30 km upstream of the head mapped "
                  "on OpenStreetMap, which is why the first point already drains 118 km². The Ramganga is given "
                  "only because it is the one river nearby with a gauge the model can be checked against.", cols=7)
        self.columns.append((ws.title, "(all rows)", "One row for each quantity, named in full; one column for "
                             "each river point."))

    def tables(self):
        t = self.t
        key = [("", [("year", "Year", INT, "Calendar year, January to December."), DAYS, WHOLE, NOTE])]
        weather = [("Weather over the catchment (ERA5)", [RAIN, PET])]
        self.add("Annual", "Year by year", t["annual"],
                 key + weather + seg_blocks(MEAN, VOLUME, PEAK, LOW7) + [BALANCE])
        key = [("", [("year", "Starts in", INT, "The calendar year in which the water year starts, on 1 June."),
                     ("water_year", "Water year", TEXT, "June to May, written 1985-86."), DAYS, WHOLE, NOTE])]
        self.add("Water year", "By water year (June to May)", t["water_year"],
                 key + weather + seg_blocks(MEAN, VOLUME, PEAK) + [BALANCE], freeze=2)
        key = [("", [("year", "Year", INT, "Calendar year."),
                     ("season", "Season", TEXT, "India Meteorological Department season."),
                     ("months", "Months", TEXT, "The months of the season."), DAYS, WHOLE, RAIN])]
        self.add("Seasons", "Season by season", t["seasons"], key + seg_blocks(MEAN, VOLUME), freeze=2)
        key = [("", [("month", "Month", MONTH, "The month, shown as year and month."),
                     ("year", "Year", INT, "Calendar year."), ("month_number", "Month number", INT, "1 = January."),
                     DAYS, WHOLE])]
        self.add("Months", "Month by month", t["months"], key + weather + seg_blocks(MEAN, VOLUME))
        daily = [("rain_mm", "Rain (mm)", ONE, RAIN[3]), ("et0_mm", "Evaporative demand (mm)", ONE, PET[3])]
        self.add("Days", "Day by day", t["days"], [
            ("", [("date", "Date", DAY, "The day, in Coordinated Universal Time (05:30 to 05:30 India Standard "
                   "Time)."), ("year", "Year", INT, "Calendar year."), ("month", "Month", INT, "1 = January."),
                  ("water_year", "Water year", TEXT, "June to May, written 1985-86."),
                  ("season", "Season", TEXT, "India Meteorological Department season.")]),
            ("Modelled daily-mean flow (m³/s)", [(f"{k}_m3s", SHORT[k], FLOW, "Mean flow of the day at the point, "
                                                  "in cubic metres a second.") for k in flow.SEGMENTS]),
            ("Weather over the catchment (ERA5)", daily)])

        avg = t["average_year_wide"].copy()
        seg = t["segments"].set_index("segment")
        total = {"name": "Whole year", "mouth_share_pct": 100.0, "rain_mm": avg["rain_mm"].sum(),
                 "et0_mm": avg["et0_mm"].sum()}
        for k in flow.SEGMENTS:
            total[f"{k}_mean_m3s"], total[f"{k}_volume_mcm"] = seg["mean_m3s"][k], avg[f"{k}_volume_mcm"].sum()
        avg = pd.concat([avg.astype({"month": "object"}), pd.DataFrame([total])], ignore_index=True)
        self.add("Average year", f"The average year, {BASE_TEXT}", avg, [
            ("", [("month", "Month number", INT, "1 = January."), ("name", "Month", TEXT, "The month.")]),
            *seg_blocks(("mean_m3s", "Mean flow (m³/s)", FLOW, f"Mean flow of the month, averaged over {BASE_TEXT}."),
                        ("volume_mcm", "Volume (million m³)", VOL, f"Volume of the month, averaged over {BASE_TEXT}.")),
            ("Whole river: spread between years", [
                ("mouth_median_m3s", "Middle year (m³/s)", FLOW, "Median of the month's mean flow over the years."),
                ("mouth_p10_m3s", "Low year (m³/s)", FLOW, "The month's mean flow is below this in one year in ten."),
                ("mouth_p90_m3s", "High year (m³/s)", FLOW, "The month's mean flow is above this in one year in ten."),
                ("mouth_share_pct", "Share of the year's volume (%)", PCT, "The month's share of the yearly volume.")]),
            ("Weather over the catchment (ERA5)", [RAIN, PET])], freeze=2)

        dep = t["dependable"].assign(segment=t["dependable"]["segment"].map(SHORT))
        self.add("Dependable", "Dependable volumes", dep, [
            ("", [("segment", "River point", TEXT, "The point on the river."),
                  ("period", "Part of the year", TEXT, "Calendar year, water year or season."),
                  ("years", "Years used", TEXT, "The years the statistic is taken over."),
                  ("n_years", "Number of years", INT, "Complete years or seasons in that span.")]),
            ("Volume (million m³)", [
                ("mean_mcm", "Mean", VOL, "Mean volume over the years."),
                ("dep50_mcm", "Met in 50 % of years", VOL, "Volume met or passed in half the years."),
                ("dep75_mcm", "Met in 75 % of years", VOL, "Volume met or passed in three years out of four: the "
                 "\"75 % dependable\" yield of Indian water planning."),
                ("dep90_mcm", "Met in 90 % of years", VOL, "Volume met or passed in nine years out of ten.")]),
            ("Extremes (million m³)", [
                ("lowest_mcm", "Lowest", VOL, "Lowest volume in the span."),
                ("lowest_year", "in", INT, "Year of the lowest volume."),
                ("highest_mcm", "Highest", VOL, "Highest volume in the span."),
                ("highest_year", "in", INT, "Year of the highest volume.")])], freeze=2, tab=GREEN)

        self.add("Flow duration", "How often a flow is reached", t["flow_duration"], [
            ("", [("years", "Years used", TEXT, "The years the curve is taken over."),
                  ("pct_of_days", "% of days", INT, "Share of days on which the flow is at least the value "
                   "shown: 90 means the river carries that much on nine days out of ten.")]),
            ("Flow equalled or exceeded on that share of days (m³/s)",
             [(f"{k}_m3s", SHORT[k], FLOW, "Flow reached or passed on that share of days.") for k in flow.SEGMENTS])],
            freeze=2, tab=GREEN)

        tr = t["trends"].assign(segment=t["trends"]["segment"].map({**SHORT, "weather": "Weather over the catchment"}))
        self.add("Trends", "Is it rising or falling?", tr, [
            ("", [("segment", "River point or weather", TEXT, "What the series belongs to."),
                  ("unit", "Unit", TEXT, "Unit of the series: m³/s for flow, mm for weather."),
                  ("series", "Yearly series", TEXT, "One value a year: the mean flow of the year or of a season, "
                   "the lowest 7-day flow, the highest daily flow, or a total of rain or evaporative demand."),
                  ("period", "Years", TEXT, "First and last year of the series."),
                  ("n", "Number of years", INT, "Years in the series."),
                  ("mean", "Mean", FLOW, "Mean of the series.")]),
            ("Theil-Sen trend", [
                ("slope_per_year", "Change a year", PVAL, "Median of the slopes between every pair of years, in the "
                 "unit of the series."),
                ("pct_per_decade", "% per decade", PCT, "The change over ten years as a share of the mean.")]),
            ("Mann-Kendall test", [
                ("tau", "Kendall's tau", TWO, "From -1 (always falling) through 0 (no order) to +1 (always rising)."),
                ("p", "p", PVAL, "Chance of so steady a run where there is no trend."),
                ("lag1", "Year-to-year persistence", TWO, "Correlation of each year with the year before, after "
                 "taking out the trend."),
                ("p_persist", "p allowing for persistence", PVAL, "The same chance after allowing for wet years "
                 "following wet years.")]),
            ("", [("reading", "Reading", TEXT, "In words, from the persistence-adjusted p: clear below 0.01, "
                   "likely below 0.05, weak sign below 0.1, otherwise no trend detected.")])],
            freeze=None, tab=GREEN)

        self.add("Decades", "Decade by decade", t["decades"], [
            ("", [("decade", "Decade", TEXT, "The years averaged."), ("years", "Years", INT, "Complete years in it.")]),
            ("Weather over the catchment (ERA5)", [RAIN, PET]),
            ("Mean yearly volume (million m³)", [(f"{k}_volume_mcm", SHORT[k], VOL, "Mean volume of a year in the "
                                                  "decade.") for k in flow.SEGMENTS]),
            ("Whole river", [
                BALANCE[1][0], BALANCE[1][1],
                ("mouth_pre_monsoon_m3s", "Mean flow, Mar-May (m³/s)", FLOW, "Mean pre-monsoon flow at the mouth."),
                ("mouth_monsoon_m3s", "Mean flow, Jun-Sep (m³/s)", FLOW, "Mean monsoon flow at the mouth."),
                ("mouth_low_7day_m3s", "Lowest 7 days (m³/s)", LOW, "Mean of the yearly lowest 7-day flows.")])],
            tab=GREEN)

        cells = ("north", "middle", "south")
        cn = self.t["thesis_years"].set_index("year")["thesis_cn2"]
        self.add("Rain and runoff", "Rain, modelled runoff and storm runoff, year by year", t["rain_runoff"], [
            ("", [("year", "Year", INT, "Calendar year."), ("days", "Days of rain data", INT, "Days of ERA5 rain "
                  "in the year."), NOTE]),
            ("ERA5 rain by grid cell (mm)", [(f"rain_{n}_mm", n.capitalize(), MM, f"Rain of the {n} cell "
                                              f"(weight {SHARE[n]:.2f}).") for n in cells]),
            ("Catchment, weighted", [RAIN, PET, ("monsoon_rain_mm", "Monsoon rain (mm)", MM, "Rain of June to "
                                                 "September.")]),
            ("Modelled river (GEOGLOWS)", [BALANCE[1][0], BALANCE[1][1]]),
            ("Storm runoff by the curve-number method (mm)", [
                (f"direct_runoff_cn{y}_mm", f"{y} land cover (CN {cn[y]})", MM, "Storm runoff of the year that "
                 f"the curve-number method gives with the thesis's {y} land cover.") for y in cn.index]),
            (f"More storm runoff, {cn.index[-1]} against {cn.index[0]}", [
                ("extra_runoff_mm", "mm", MM, "Difference between the two land covers for the same rain."),
                ("extra_runoff_mcm", "million m³", VOL, f"The same over the {MOUTH_KM2:g} km² catchment.")])],
            tab=GREEN)

    def scenario(self):
        t = self.t
        ws = self.sheet("Land-cover scenario", "The same rain on more built-up land", tab=GREEN)
        sc, ty = t["scenario"].set_index("cn2"), t["thesis_years"].set_index("year")
        cn79, cn09 = ty["thesis_cn2"][1979], ty["thesis_cn2"][2009]
        ws.cell(5, 1, f"A. Storm runoff of an average year ({BASE_TEXT} rain) for a range of curve numbers").font = SECTION
        r = self.table(ws, t["scenario"], [
            ("Curve number (CN)", [
                ("cn2", "Average soil (CN-II)", TWO, "Area-weighted curve number for average soil moisture."),
                ("land_cover", "Which land cover", TEXT, "The three values found in the thesis are named."),
                ("cn1_dry", "Dry soil (CN-I)", TWO, "Used when under 35.6 mm fell in the five days before."),
                ("cn3_wet", "Wet soil (CN-III)", TWO, "Used when over 53.3 mm fell in the five days before.")]),
            ("Storm runoff of an average year", [
                ("direct_runoff_mm", "mm", MM, "Mean yearly storm runoff, worked out day by day."),
                ("pct_of_rain", "% of rain", PCT, "Storm runoff as a share of the rain."),
                ("volume_mcm", "million m³", VOL, f"The same over the {MOUTH_KM2:g} km² catchment.")]),
            (f"Change from 1979 (CN {cn79})", [
                ("change_mm", "mm", MM, "Difference in storm runoff from the 1979 curve number."),
                ("change_pct", "%", PCT, "The same as a share of the 1979 runoff."),
                ("change_mcm", "million m³", VOL, "The same as a volume over the catchment.")])], top=6)
        ws.cell(r, 1, "B. The three years of the thesis, set beside the same method on ERA5 rain").font = SECTION
        r = self.table(ws, t["thesis_years"], [
            ("", [("year", "Year", INT, "The year of the land-cover map.")]),
            ("From the thesis (its 4,120 km² study area)", [
                ("thesis_cn2", "Curve number (CN-II)", TWO, "Area-weighted curve number of the study area."),
                ("thesis_rain_mm", "Rain of the year (mm)", MM, "Rain the thesis used for that year."),
                ("thesis_runoff_mm", "Runoff (mm)", MM, "Runoff the thesis worked out for that year."),
                ("thesis_recharge_mm", "Recharge (mm)", MM, "Groundwater recharge the thesis worked out.")]),
            ("Same method, ERA5 rain over the Nakatiya catchment", [
                ("era5_rain_mm", "Rain of the year (mm)", MM, "ERA5 rain of that year, weighted over the catchment."),
                ("era5_south_rain_mm", "Rain, south cell (mm)", MM, "ERA5 rain of the cell that holds Bareilly city."),
                ("era5_runoff_own_cn_mm", "Runoff, that year's CN (mm)", MM, "Storm runoff with that year's curve "
                 "number."),
                ("era5_runoff_cn1979_mm", f"Runoff, CN {cn79} (mm)", MM, "Storm runoff had the land stayed as in "
                 "1979."),
                ("era5_runoff_cn2009_mm", f"Runoff, CN {cn09} (mm)", MM, "Storm runoff had the land been as in "
                 "2009.")])], top=r + 1)
        ws.cell(r, 1, "How to read this").font = SECTION
        notes = [
            "The method is the Soil Conservation Service curve-number method, applied day by day exactly as in the "
            "thesis: storm runoff Q = (P - 0.2 S)² / (P + 0.8 S), with P the day's rain and S = 25400 / CN - 254 "
            "(both in mm); no runoff until the rain exceeds 0.2 S; and the rain of the five days before decides "
            "whether the dry, average or wet curve number applies. It reproduces the thesis's worked day (3 July "
            "1977: 28.4 mm of rain gives 7.96 mm of runoff).",
            f"Table A holds the rain fixed and changes only the land. Going from the 1979 curve number ({cn79}) to "
            f"the 2009 one ({cn09}) raises storm runoff from {sc['direct_runoff_mm'][cn79]:.0f} to "
            f"{sc['direct_runoff_mm'][cn09]:.0f} mm a year: {sc['change_pct'][cn09]:+.0f} %, or "
            f"{sc['change_mcm'][cn09]:.1f} million m³ a year over the Nakatiya catchment. This water leaves as a "
            "quick flush during storms instead of soaking into the ground, which is the first half of the argument "
            "that building over the land starves the river in the dry season.",
            f"Table B explains why the thesis's own runoff figures rise much more steeply "
            f"({ty['thesis_runoff_mm'][1979]:.0f}, {ty['thesis_runoff_mm'][1990]:.0f} and "
            f"{ty['thesis_runoff_mm'][2009]:.0f} mm): it compares three different years, and 1979 was a drought "
            f"year ({ty['thesis_rain_mm'][1979]:.0f} mm of rain). Most of that rise is the difference in rain "
            "between the years, not the change in land cover. The two must be separated before the numbers are "
            "used as evidence, and table A does that.",
            "Limits: the curve numbers describe the thesis's whole study area, not the Nakatiya catchment, whose own "
            "land-cover history is still to be worked out. The method counts storm runoff only; it says nothing "
            "about seepage from the ground into the river, and it cannot tell land near the bank from land far "
            "from it. The thesis's rain came from rain gauges and ERA5 rain differs from it year by year, so the "
            "right-hand block of table B shows the method at work, not a reconstruction of the thesis.",
        ]
        for text in notes:
            r = self.note(ws, r + 1, text, cols=10) - 1
        return ws

    def width(self):
        t = self.t
        w = t["satellite_width"].assign(reach=t["satellite_width"]["reach"].map(REACH))
        ws = self.add("Satellite width", "Open water seen from space", w, [
            ("", [("reach", "Reach", TEXT, "The stretch of river: a corridor around the mapped channel."),
                  ("date", "Date of the image", DAY, "Day of the Sentinel-2 image."),
                  ("period", "Time of year", TEXT, "After the monsoon (October-November) or in the dry months."),
                  ("width_m", "Open-water width (m)", TWO, "Open-water area of the reach divided by its length. "
                   f"Land away from the river gives up to about {WIDTH_NOISE_M} m by the same sum, so smaller "
                   "values mean no open water was seen."),
                  ("reach_km", "Length of the reach (km)", TWO, "River length inside the reach."),
                  ("area_ha", "Open water (hectares)", TWO, "Width x length."),
                  ("reliable", "Reliable?", TEXT, "\"no\" where the colour of pure water had to be borrowed from "
                   "another date or was mixed with land: the width is then uncertain."),
                  ("water_signature", "Where the colour of water came from", TEXT, "The same day, another date, or "
                   "the same day but mixed with land."),
                  ("mouth_flow_m3s", "Modelled flow at the mouth that day (m³/s)", FLOW, "For comparison only: the "
                   "model does not see what the satellite sees."),
                  ("scene", "Sentinel-2 image", TEXT, "Identifier of the image.")])], freeze=2, tab=GREEN)
        self.widths[ws.title, 10] = 30
        self.note(ws, ws.max_row + 2, "A width is an area of open water, not a flow. It cannot be turned into a flow "
                  "without a rating measured in the field (the same place gauged at several water levels). Widths "
                  "come from working out the share of water in each 10 m pixel (sub-pixel unmixing), because the "
                  "river is narrower than one pixel; floating weed, shadow and wet sand all blur the answer. The "
                  "free satellites cannot see water under the city's buildings and trees at all.", cols=9)

    def checks(self):
        t = self.t
        ws = self.sheet("Checks", "What the model can be checked against", tab=GREEN)
        a, b = flow.CHAUBARI_PERIOD
        ws.cell(5, 1, f"A. Against a gauge-calibrated model of the Ramganga at Chaubari, {a} to {b}").font = SECTION
        r = self.table(ws, t["gauge_check"], [
            ("", [("season", "Season", TEXT, "Monsoon or pre-monsoon."), ("months", "Months", TEXT, "Its months."),
                  ("met_in_pct_of_years", "Met in % of years", INT, "The seasonal mean flow is at least the value "
                   "shown in this share of years.")]),
            ("Seasonal mean flow (m³/s)", [
                ("geoglows_m3s", "This model (GEOGLOWS)", VOL, "From the daily flows of this workbook."),
                ("swat_present_m3s", "Calibrated model, today's river", VOL, "WWF-India / INRM model calibrated to "
                 "the Central Water Commission gauge, with today's dams, canals and irrigation."),
                ("swat_natural_m3s", "Calibrated model, natural river", VOL, "The same model without dams, canals "
                 "or irrigation.")]),
            ("", [("ratio_to_present", "This model ÷ today's river", TWO, "How many times higher this model is.")])],
            top=6)
        ratio = t["gauge_check"].groupby("season")["ratio_to_present"].agg(["min", "max"])
        gauge, model = flow.CHAUBARI_GAUGE_AREA_KM2, flow.SEGMENTS["ramganga"]["area_km2"]
        r = self.note(ws, r, "Reading: where it can be checked, this model gives "
                      f"{ratio['min']['monsoon']:.1f} to {ratio['max']['monsoon']:.1f} times the monsoon flow and "
                      f"{ratio['min']['pre-monsoon']:.1f} to {ratio['max']['pre-monsoon']:.1f} times the pre-monsoon "
                      "flow of a model calibrated to the gauge. Water taken out upstream (the dam at Kalagarh, "
                      "canals, irrigation) explains part of the gap and model error the rest; the two cannot be "
                      "separated here. The ratios must not be used to scale the Nakatiya: the Ramganga is a "
                      "Himalayan river with a dam, the Nakatiya a stream of the plain. Catchment area at the site: "
                      f"{gauge:,} km² by the Central Water Commission, {model:,.0f} km² in this model "
                      f"({100 * (model / gauge - 1):.0f} % more).", cols=7)

        ws.cell(r + 1, 1, f"B. Four ways to estimate how much of the rain reaches the river (average year {BASE_TEXT})"
                ).font = SECTION
        r = self.table(ws, t["yield_estimates"], [
            ("", [("method", "Method", TEXT, "How the estimate is made."),
                  ("what_it_counts", "What it counts", TEXT, "Which water the estimate includes."),
                  ("runoff_mm", "Runoff (mm a year)", MM, "Depth of water over the catchment."),
                  ("pct_of_rain", "% of rain", PCT, "Share of the rain."),
                  ("volume_mcm", "million m³ a year", VOL, f"Over the {MOUTH_KM2:g} km² catchment."),
                  ("mean_m3s", "As a steady flow (m³/s)", FLOW, "The same volume as a constant flow.")])], top=r + 2)
        y = t["yield_estimates"]
        r = self.note(ws, r, f"Reading: the estimates span {y['runoff_mm'].min():.0f} to {y['runoff_mm'].max():.0f} "
                      "mm a year. The curve-number figures count storm runoff only, so they should be the lowest. "
                      "The water-balance curve is known to run low where nearly all the rain falls in one season. "
                      "A fair statement is that the weather puts something between a sixth and a third of the rain "
                      "into the river, and that the true figure is not known to better than a factor of about two "
                      "until the river is measured.", cols=7)

        ws.cell(r + 1, 1, "C. Were the first two years a start-up artefact?").font = SECTION
        r = self.table(ws, t["dry_years"], [
            ("", [("year", "Year", TEXT, "The year or span."), ("what", "What it is", TEXT, "Why the row is here."),
                  ("rain_mm", "Rain (mm)", MM, RAIN[3]),
                  ("mouth_runoff_mm", "Runoff depth (mm)", MM, BALANCE[1][0][3]),
                  ("pct_of_rain", "Runoff as % of rain", PCT, "Runoff depth divided by rain."),
                  ("mouth_mean_m3s", "Mean flow (m³/s)", FLOW, "Mean flow at the mouth."),
                  ("mouth_low_7day_m3s", "Lowest 7 days (m³/s)", LOW, LOW7[3]),
                  ("zero_flow_days", "Days with no flow", ONE, "Days on which the model gives under 0.01 m³/s.")])],
            top=r + 2)
        d = t["dry_years"].set_index("year")
        ann = t["annual"][t["annual"]["complete"]]
        kept, every = (ann.loc[ann["year"] >= y0, "mouth_volume_mcm"].mean() for y0 in (FIRST, ann["year"].min()))
        r = self.note(ws, r, f"Reading: in 1940 and 1941 the model turned {d['pct_of_rain']['1940']:.1f} and "
                      f"{d['pct_of_rain']['1941']:.1f} % of the rain into river flow, the same as in the 1987 "
                      f"drought ({d['pct_of_rain']['1987']:.1f} %), and the river ran dry for some days in the "
                      "year after each. Both were dry years in ERA5 "
                      f"({d['rain_mm']['1940']:.0f} and {d['rain_mm']['1941']:.0f} mm of rain against "
                      f"{d['rain_mm'].iloc[-1]:,.0f} mm in an average year). Nothing marks them as a start-up "
                      "artefact. They are still left out of the statistics, because GEOGLOWS does not say how the "
                      f"run was started; putting them back lowers the mean yearly volume of the record by "
                      f"{100 * (1 - every / kept):.1f} %.", cols=7)

    def calculators(self):
        ws = self.sheet("Calculators", "Work out a flow or a runoff yourself", tab=AMBER, banner=False)
        ws["A3"] = ("Type your own numbers into the yellow cells. Everything in bold is a formula; the grey column "
                    "shows what the sample numbers should give.")
        ws["A3"].font = BOLD
        self.widths[ws.title, 1] = 60
        for col in range(2, 12):
            self.widths[ws.title, col] = 12 if col < 5 else 8
        self.widths[ws.title, 4] = 17

        p = Panel(ws, 5)
        p.heading("A. Float method: discharge from a tape, a float and a stopwatch",
                  "Pick a straight stretch 10 to 20 m long. Measure the depth at equal steps across the water "
                  "(not at the banks) and time a half-sunk float at least three times.")
        p.enter("width", "Width of the water, bank to bank", 4, "m")
        p.enter_many("depths", "Depths at equal steps across the water (m), up to ten", [0.2, 0.4, 0.2])
        p.enter("distance", "Distance the float travels", 10, "m")
        p.enter_many("times", "Time of each run (seconds), up to ten", [19, 20, 21])
        p.enter("coeff", "Surface coefficient (0.8 rough and shallow, 0.9 smooth and deep)", 0.85)
        p.show("n", "Number of depth readings", "COUNT({depths})", "", 3, INT)
        p.show("step", "Step between readings", "IF({n}>0,{width}/({n}+1),0)", "m", 1, TWO)
        p.show("area", "Cross-section of the water", "{step}*SUM({depths})", "m²", 0.8, TWO)
        p.show("surface", "Speed at the surface", "IF(COUNT({times})>0,{distance}/AVERAGE({times}),0)", "m/s", 0.5, TWO)
        p.show("q", "Discharge", "{area}*{surface}*{coeff}", "m³/s", 0.34, LOW)
        p.show("q_low", "   if the coefficient were 0.8", "{area}*{surface}*0.8", "m³/s", 0.32, LOW)
        p.show("q_high", "   if the coefficient were 0.9", "{area}*{surface}*0.9", "m³/s", 0.36, LOW)
        p.show("lps", "Discharge in litres a second", "{q}*1000", "L/s", 340, MM)
        p.show("mld", "Discharge in million litres a day", "{q}*86.4", "million L/day", 29.376, ONE)
        p.show("cusec", "Discharge in cusec (cubic feet a second)", "{q}*35.3147", "cusec", 12.007, ONE)

        p = Panel(ws, p.row + 1)
        p.heading("B. Manning's formula: discharge from the shape and slope of the channel",
                  "For a channel with a flat bed and sloping sides. Useful for a flood mark: measure the channel "
                  "up to the mark when the water has gone.")
        p.enter("b", "Width of the bed", 6, "m")
        p.enter("y", "Depth of water", 0.5, "m")
        p.enter("z", "Side slope: metres sideways for each metre up (0 for upright banks)", 1.5)
        p.enter("s", "Slope of the water surface: fall divided by length (4 cm in 100 m = 0.0004)", 0.0004)
        p.enter("n", "Roughness n of the bed (see the table below)", 0.035)
        p.show("area", "Cross-section of the water, (b + z y) y", "({b}+{z}*{y})*{y}", "m²", 3.375, LOW)
        p.show("wet", "Wetted perimeter, b + 2 y √(1 + z²)", "{b}+2*{y}*SQRT(1+{z}^2)", "m", 7.8028, LOW)
        p.show("r", "Hydraulic radius R, area ÷ wetted perimeter", "{area}/{wet}", "m", 0.4325, LOW)
        p.show("v", "Mean speed, R^(2/3) √S ÷ n", "{r}^(2/3)*SQRT({s})/{n}", "m/s", 0.3268, LOW)
        p.show("q", "Discharge", "{area}*{v}", "m³/s", 1.103, LOW)
        p.show("mld", "Discharge in million litres a day", "{q}*86.4", "million L/day", 95.3, ONE)
        p.line("Roughness n (Chow, 1959): usual value, and range", BOLD)
        for label, usual, span in (
                ("Natural stream, clean and straight", 0.030, "0.025 to 0.033"),
                ("The same with more stones and weeds", 0.035, "0.030 to 0.040"),
                ("Clean but winding, with pools and shoals", 0.040, "0.033 to 0.045"),
                ("Sluggish, weedy, with deep pools", 0.070, "0.050 to 0.080"),
                ("Very weedy reaches", 0.100, "0.075 to 0.150"),
                ("Dug earth channel, clean, after weathering", 0.022, "0.018 to 0.025"),
                ("Dug earth channel with short grass", 0.027, "0.022 to 0.033"),
                ("Concrete, float finish", 0.015, "0.013 to 0.016")):
            ws.cell(p.row, 1, "   " + label)
            ws.cell(p.row, 2, usual).number_format = LOW
            ws.cell(p.row, 3, span)
            p.row += 1

        p = Panel(ws, p.row + 1)
        p.heading("C. Storm runoff of one day by the curve-number method",
                  "The sample is the thesis's worked day, 3 July 1977, with the 1979 land cover.")
        p.enter("p", "Rain of the day", 28.4, "mm")
        p.enter("before", "Rain of the five days before", 63.4, "mm")
        p.enter("cn", "Curve number for average soil moisture (CN-II)", 75.63)
        p.enter("km2", "Catchment area", MOUTH_KM2, "km²")
        dry, wet = flow.AMC_LIMITS_MM
        p.show("amc", "Soil moisture class", f'IF({{before}}<{dry},"I (dry)",IF({{before}}>{wet},"III (wet)",'
               '"II (average)"))', "", "III (wet)", TEXT)
        p.show("used", "Curve number used", f"IF({{before}}<{dry},4.2*{{cn}}/(10-0.058*{{cn}}),"
               f"IF({{before}}>{wet},23*{{cn}}/(10+0.13*{{cn}}),{{cn}}))", "", 87.71, TWO)
        p.show("s", "Retention S, 25400 ÷ CN - 254", "25400/{used}-254", "mm", 35.59, TWO)
        p.show("q", "Storm runoff Q, (P - 0.2 S)² ÷ (P + 0.8 S)", "IF({p}>0.2*{s},({p}-0.2*{s})^2/({p}+0.8*{s}),0)",
               "mm", 7.97, TWO)
        p.line("The thesis's table gives 7.96 mm for this day: it rounds the curve number and S before the last step.",
               QUIET)
        p.show("share", "Share of the rain that runs off", "IF({p}>0,100*{q}/{p},0)", "%", 28.0, ONE)
        p.show("volume", "Runoff volume over the catchment", "{q}*{km2}/1000", "million m³", 2.96, TWO)

        p = Panel(ws, p.row + 1)
        p.heading("D. From a flow to a volume, and between units", "Change the yellow cells; the rest follows.")
        p.enter("q", "Flow", 1, "m³/s")
        p.enter("days", "Kept up for", 365, "days")
        p.enter("km2", "Catchment area", MOUTH_KM2, "km²")
        p.show("lps", "Flow in litres a second", "{q}*1000", "L/s", 1000, MM)
        p.show("cusec", "Flow in cusec (cubic feet a second)", "{q}*35.3147", "cusec", 35.31, TWO)
        p.show("mld", "Flow in million litres a day", "{q}*86.4", "million L/day", 86.4, ONE)
        p.show("ham", "Flow in hectare-metres a day", "{q}*8.64", "ha-m/day", 8.64, TWO)
        p.show("vol", "Volume over those days", "{q}*{days}*0.0864", "million m³", 31.536, LOW)
        p.show("depth", "The same as a depth over the catchment", "1000*{vol}/{km2}", "mm",
               round(31536 / MOUTH_KM2, 1), ONE)
        p.show("v_ham", "The volume in hectare-metres", "{vol}*100", "ha-m", 3153.6, VOL)
        p.show("v_crore", "The volume in crore litres", "{vol}*100", "crore L", 3153.6, VOL)
        p.show("v_af", "The volume in acre-feet", "{vol}*810.713", "acre-feet", 25566.6, VOL)
        p.show("v_tmc", "The volume in TMC (thousand million cubic feet)", "{vol}/28.3168", "TMC", 1.1137, PVAL)
        ws.freeze_panes = "B5"

    def limits(self):
        t = self.t
        ratio = t["gauge_check"].groupby("season")["ratio_to_present"].agg(["min", "max"])
        ann = t["annual"][t["annual"]["complete"]]
        kept, every = (ann.loc[ann["year"] >= y0, "mouth_volume_mcm"].mean() for y0 in (FIRST, ann["year"].min()))
        points = [
            "Modelled, not measured. There is no gauge on the Nakatiya, and nothing here has been checked against "
            "a measurement on this river. Treat every flow as right to within a factor of about two at best, and "
            "worse at low flow.",
            "Weather only. The model has no city, sewage, drains, pumping, irrigation, canals, dams or aquifer. The "
            "real river through Bareilly receives sewage and drain water the model does not know about, and loses "
            "water it does not know about.",
            "Coarse weather. ERA5 works on cells about 30 km across; the whole Nakatiya catchment "
            f"({MOUTH_KM2:g} km²) is smaller than one cell. A cloudburst over part of the catchment is smoothed "
            "away, so flood peaks are understated and their days approximate.",
            "The one check available says the model runs high. On the Ramganga at Chaubari it gives "
            f"{ratio['min']['monsoon']:.1f} to {ratio['max']['monsoon']:.1f} times the monsoon flow and "
            f"{ratio['min']['pre-monsoon']:.1f} to {ratio['max']['pre-monsoon']:.1f} times the pre-monsoon flow of "
            "a gauge-calibrated model of today's river (Checks sheet). That ratio must not be used to scale the "
            "Nakatiya: the two rivers are of different kinds.",
            "Flat-land catchments are uncertain. The model's river network was drawn from satellite elevation "
            "data. On the nearly flat Bareilly plain, roads, canals and embankments turn water in ways it cannot "
            "see, so the catchment areas, and with them the volumes, may be off. The model's river also starts 25 "
            "to 30 km upstream of the head mapped on OpenStreetMap.",
            "Early decades. ERA5 before about 1960 rests on far fewer observations. The fall in rain and flow from "
            "the 1950s to the 1980s is in ERA5; it must be checked against the rain grids of the India "
            "Meteorological Department before it is quoted as fact.",
            "The first two years, 1940 and 1941, are left out of every statistic as a warm-up precaution. The "
            "Checks sheet shows they behave like any drought year; putting them back lowers the mean yearly volume "
            f"of the record by {100 * (1 - every / kept):.1f} %.",
            "Low flows are the least certain part. ERA5's land model has no groundwater store: its dry-season flow "
            "is drainage from a soil column 2.9 m deep. The real river's dry-season flow depends on the water "
            "table, which pumping has lowered. Flows are stored to 0.01 m³/s, so values under about 0.05 m³/s "
            "carry one significant figure.",
            "A rising modelled dry-season flow is a statement about the weather of the reconstruction (chiefly a "
            "fall in its evaporative demand), not about the river. It says what the weather would have done; it "
            "cannot show what the city has done.",
            "A \"highest day\" is the highest daily mean, not the flood peak, which on a small river can be "
            "several times higher and last a few hours.",
            "Dependable volumes (50, 75 and 90 %) are statistics of the model. They must not be used to design or "
            "allocate anything until the model has been calibrated against measurements.",
            "The land-cover scenario counts storm runoff only, uses the curve numbers of the thesis's whole "
            "4,120 km² study area, and cannot represent seepage near the bank. It answers one question: how much "
            "more storm runoff does the same rain give.",
            "A satellite width is an area of open water, not a flow. Rows marked \"no\" under \"Reliable?\" are "
            "uncertain, and no flow series has been derived from the widths.",
            "Days do not line up exactly: flow days are days of Coordinated Universal Time, rain days are India "
            "Standard Time days, so a storm can appear a day apart in the two.",
            "A trend test says whether a series rises or falls more steadily than chance would give. It says "
            "nothing about the cause.",
            "None of this is a legal or regulatory finding. The figures are flags for where to measure.",
        ]
        self.text_sheet("Limits", "What these numbers cannot tell you",
                        [("p", f"{i}.  {text}") for i, text in enumerate(points, 1)])

    def dictionary(self):
        ws = self.sheet("Dictionary", "Every column explained", tab=SLATE, banner=False)
        seen, rows = {}, []
        for sheet, label, meaning in self.columns:
            if (sheet, label) in seen:
                seen[sheet, label][3] = "yes"
                continue
            seen[sheet, label] = row = [sheet, label, meaning, ""]
            rows.append(row)
        cols = pd.DataFrame(rows, columns=["sheet", "column", "meaning", "repeated"])
        r = self.table(ws, cols, [("", [("sheet", "Sheet", TEXT, ""), ("column", "Column", TEXT, ""),
                                        ("meaning", "What it means", TEXT, ""),
                                        ("repeated", "One for each river point?", TEXT, "")])],
                       top=4, freeze=0, filtered=True, describe=False)
        self.widths[ws.title, 3] = 110
        for row in range(5, r - 1):
            ws.cell(row, 3).alignment = WRAP
        ws.cell(r, 1, "Terms").font = SECTION
        for i, (term, meaning) in enumerate(TERMS, r + 1):
            ws.cell(i, 1, term).alignment = WRAP
            ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
            ws.cell(i, 2, meaning).alignment = WRAP
            ws.row_dimensions[i].height = 15 * math.ceil(len(meaning) / 140) + 4
        self.widths[ws.title, 1] = 24
        self.widths[ws.title, 2] = 38

    def charts(self, ws):
        t = self.t
        cd = self.sheet("Chart data", "The numbers behind the charts", tab=SLATE, banner=False)
        head, points = 4, [SHORT[k] for k in flow.NAKATIYA_KEYS]

        def put(col, labels, rows):
            for j, label in enumerate(labels):
                cell = cd.cell(head, col + j, label)
                cell.font, cell.fill, cell.alignment = HEAD, FILL_HEAD, CENTRE
                self.widen(cd, col + j, 15)
            for i, row in enumerate(rows, head + 1):
                for j, v in enumerate(row):
                    if clean(v) is not None:
                        cd.cell(i, col + j, clean(v))
            return head + len(rows)

        def ref(col, last, header=True):
            return Reference(cd, min_col=col, min_row=head + (not header), max_row=last)

        def frame(chart, title, x_title, y_title, legend=True):
            chart.title, chart.x_axis.title, chart.y_axis.title = title, x_title, y_title
            chart.width, chart.height = 16.5, 8.6
            chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl would otherwise hide both axes
            chart.roundedCorners = False
            for title in (chart.title, chart.x_axis.title, chart.y_axis.title):
                title.overlay = False  # otherwise Excel draws titles over the plot
            for axis in (chart.x_axis, chart.y_axis):
                if axis.majorGridlines is not None:
                    axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9"))
            if legend:
                chart.legend.position = "b"
                chart.legend.overlay = False
            else:
                chart.legend = None
            return chart

        def stroke(series, colour, width=2.25, marker=None):
            series.graphicalProperties.line.solidFill = colour
            series.graphicalProperties.line.width = int(width * 12700)
            series.smooth = False
            series.marker.symbol = marker or "none"
            if marker:
                series.marker.size = 5
                series.marker.graphicalProperties.solidFill = colour
                series.marker.graphicalProperties.line.solidFill = colour

        def dots(values, xs, colour):
            series = Series(values, xs, title_from_data=True)
            stroke(series, colour, marker="circle")
            series.graphicalProperties.line.solidFill = None  # a line holds a fill or noFill, never both
            series.graphicalProperties.line.noFill = True
            return series

        cd.row_dimensions[head].height = 48
        ann = t["annual"][t["annual"]["complete"]]
        kept = ann[ann["year"] >= FIRST]
        m = t["segments"].set_index("segment").loc["mouth"]
        sc, cn = t["scenario"], t["thesis_years"].set_index("year")["thesis_cn2"]

        last = put(1, ["Year", "Volume of the year (million m³)", "Mean of the last 10 years (million m³)"],
                   list(zip(ann["year"].tolist(), ann["mouth_volume_mcm"].tolist(),
                            ann["mouth_volume_10yr_mcm"].tolist())))
        c1, mean10 = BarChart(), LineChart()
        c1.add_data(ref(2, last), titles_from_data=True)
        c1.set_categories(ref(1, last, header=False))
        mean10.add_data(ref(3, last), titles_from_data=True)
        c1.series[0].graphicalProperties.solidFill = BLUES[0]
        c1.series[0].graphicalProperties.line.solidFill = BLUES[0]
        stroke(mean10.series[0], NAVY)
        c1.gapWidth = 40
        c1.x_axis.tickLblSkip = c1.x_axis.tickMarkSkip = 10
        c1 += mean10
        frame(c1, "Water carried by the Nakatiya each year (modelled)", "Year", "million m³")

        avg = t["average_year_wide"]
        last = put(5, ["Month", *points], [(name[:3], *[avg[f"{k}_mean_m3s"][i] for k in flow.NAKATIYA_KEYS])
                                           for i, name in enumerate(avg["name"])])
        c2 = LineChart()
        for j, colour in enumerate(BLUES):
            c2.add_data(ref(6 + j, last), titles_from_data=True)
            stroke(c2.series[j], colour)
        c2.set_categories(ref(5, last, header=False))
        frame(c2, f"The average year, {BASE_TEXT}", "Month", "Mean flow (m³/s)")

        s = t["seasons"]
        dry = s[(s["season"] == "pre-monsoon") & s["complete"] & (s["year"] >= FIRST)].set_index("year")["mouth_mean_m3s"]
        fit = flow.trend(dry.loc[flow.TREND_START:])
        fitted = [fit["intercept"] + fit["slope_per_year"] * y if y >= flow.TREND_START else None for y in dry.index]
        last = put(11, ["Year", "Mean flow, March to May (m³/s)", f"Trend since {flow.TREND_START} (m³/s)"],
                   list(zip(dry.index.tolist(), dry.tolist(), fitted)))
        c3 = ScatterChart()
        c3.scatterStyle = "lineMarker"
        xs = ref(11, last, header=False)
        line = Series(ref(13, last), xs, title_from_data=True)
        stroke(line, RED)
        c3.series += [dots(ref(12, last), xs, BLUES[1]), line]
        c3.x_axis.scaling.min, c3.x_axis.scaling.max, c3.x_axis.majorUnit = 1940, 2030, 10
        c3.x_axis.number_format = c3.y_axis.number_format = "General"
        frame(c3, "Dry-season flow at the mouth, March to May (modelled)", "Year", "Mean flow (m³/s)")

        f = t["flow_duration"]
        f = f[f["years"] == f"{BASE[0]}-{BASE[1]}"]
        last = put(15, ["% of days", *points],
                   list(zip(f["pct_of_days"].tolist(), *[f[f"{k}_m3s"].tolist() for k in flow.NAKATIYA_KEYS])))
        c4 = ScatterChart()
        c4.scatterStyle = "lineMarker"
        xs = ref(15, last, header=False)
        for j, colour in enumerate(BLUES):
            curve = Series(ref(16 + j, last), xs, title_from_data=True)
            stroke(curve, colour)
            c4.series.append(curve)
        c4.y_axis.scaling.logBase = 10
        c4.y_axis.scaling.min, c4.y_axis.scaling.max = 0.01, 100
        c4.x_axis.scaling.min, c4.x_axis.scaling.max, c4.x_axis.majorUnit = 0, 100, 10
        c4.x_axis.crosses = "min"
        c4.x_axis.number_format = c4.y_axis.number_format = "General"
        frame(c4, f"How often a flow is reached, {BASE_TEXT}", "% of days the flow is equalled or exceeded",
              "Flow (m³/s), log scale")

        last = put(21, ["Rain of the year (mm)", "Runoff depth of the year (mm)"],
                   list(zip(kept["rain_mm"].tolist(), kept["mouth_runoff_mm"].tolist())))
        c5 = ScatterChart()
        c5.scatterStyle = "lineMarker"
        c5.series.append(dots(ref(22, last), ref(21, last, header=False), BLUES[2]))
        c5.x_axis.number_format = c5.y_axis.number_format = "General"
        frame(c5, f"Rain against modelled runoff, {FIRST} to {int(kept['year'].max())}", "Rain of the year (mm)",
              "Runoff depth (mm)", legend=False)

        last = put(24, ["Curve number (CN-II)", "Storm runoff (mm a year)"],
                   list(zip(sc["cn2"].tolist(), sc["direct_runoff_mm"].tolist())))
        c6 = ScatterChart()
        c6.scatterStyle = "lineMarker"
        curve = Series(ref(25, last), ref(24, last, header=False), title_from_data=True)
        stroke(curve, NAVY, marker="circle")
        c6.series.append(curve)
        c6.x_axis.scaling.min, c6.x_axis.scaling.max, c6.x_axis.majorUnit = 68, 92, 2
        c6.x_axis.number_format = c6.y_axis.number_format = "General"
        frame(c6, f"Storm runoff for the same {BASE_TEXT} rain as land is built over", "Curve number (CN-II)",
              "Storm runoff (mm a year)", legend=False)

        change = sc.set_index("cn2")["change_pct"][cn[2009]]
        captions = [
            "Volume of each year at the mouth (bars) and the mean of the last ten years (line). The 1950s and "
            f"1960s were wet and the 1980s dry; since {flow.TREND_START} the yearly flow shows "
            f"{say(self.trend('mouth', 'year'))}.",
            f"Mean flow of each month at the four points. {m['monsoon_share_pct']:.0f} % of the year's water passes "
            "from June to September, and May is the lowest month.",
            f"One dot a year, with the trend since {flow.TREND_START}: "
            f"{say(self.trend('mouth', 'pre-monsoon (Mar-May)'))}. The model knows only the weather, so this is "
            "what rain and evaporation alone would have given.",
            "The flow that is equalled or exceeded on a given share of days. The vertical scale is logarithmic: "
            "each step up is ten times more water.",
            f"One dot a year. A wet year gives far more than its share: correlation "
            f"{kept['rain_mm'].corr(kept['mouth_runoff_mm']):.2f}, and runoff roughly triples between a 900 mm and "
            "a 1,400 mm year.",
            f"Storm runoff of an average year by the curve-number method. The change found in the thesis, from "
            f"{cn[1979]} (1979) to {cn[2009]} (2009), adds {change:.0f} %.",
        ]
        for i, (chart, caption) in enumerate(zip((c1, c2, c3, c4, c5, c6), captions)):
            row, col = 5 + 20 * (i // 2), 1 + 10 * (i % 2)
            ws.add_chart(chart, f"{get_column_letter(col)}{row}")
            ws.merge_cells(start_row=row + 17, start_column=col, end_row=row + 19, end_column=col + 8)
            cell = ws.cell(row + 17, col, caption)
            cell.alignment, cell.font = WRAP, QUIET


def write_workbook(t, path):
    """Write the tables of build_flow_history.build_tables() to `path` as a workbook for reading."""
    w = Writer(t)
    w.read_me()
    charts = w.sheet("Charts", "The main results in six pictures")
    w.points()
    w.tables()
    w.scenario()
    w.width()
    w.checks()
    w.calculators()
    w.limits()
    w.dictionary()
    w.charts(charts)
    w.finish()
    w.wb.properties.title = "Nakatiya river: modelled flow and volume"
    w.wb.properties.creator = "KhetOS river observatory"
    w.wb.properties.description = BANNER
    w.wb.save(path)
